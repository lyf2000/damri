from io import BytesIO
from typing import Iterable, Tuple, TypeVar

from damri.models.base import BaseModel
from django.template.loader import render_to_string
from openpyxl import Workbook


class HeadingColumn(BaseModel):
    name: str


class ItemRow(BaseModel):
    pass


ItemRowType = TypeVar("ItemRowType", bound=ItemRow)


class BaseExcelGeneratorService:
    HEADINGS: Tuple[HeadingColumn, ...] = tuple()

    def __init__(self):
        self.wb = Workbook()
        self.ws = self.wb.active

    def generate(self) -> BytesIO:
        self._add_headings()
        self._fill()
        return self.save()

    def save(self) -> BytesIO:
        buff = BytesIO()
        self.wb.save(buff)
        return buff

    def _fill(self) -> None:
        for row in self._iterate_rows():
            self._append_row(row.values())

    def _add_headings(self):
        self._append_row([col.name for col in self.HEADINGS])

    def _append_row(self, row):
        self.ws.append([str(col) for col in row])

    def _iterate_rows(self) -> Iterable[ItemRowType]:
        raise NotImplementedError

    def _iterate_cells(self, x0, y0, x1=None, y1=None) -> Iterable:
        x1 = x1 or x0
        y1 = y1 or y0

        for cell in self.ws[f"{x0}{y0}":f"{x1}{y1}"]:
            yield cell


class ExcelToHTMLTemplateGeneratorMixin(BaseExcelGeneratorService):
    """
    Класс-обертка для генерации хтмл-файла из табличных данных
    """

    TEMPLATE_PATH: str = ""

    @property
    def template(self) -> str:
        return self.TEMPLATE_PATH

    def generate(self) -> BytesIO:
        context = self.get_context()
        return self.save(context)

    def save(self, context: dict) -> BytesIO:
        page = str(render_to_string(self.template, context).replace("\n", ""))
        return BytesIO(page.encode())

    def get_context(self) -> dict:
        return dict(
            headings=self.HEADINGS,
            rows=self._iterate_rows(),
        )
